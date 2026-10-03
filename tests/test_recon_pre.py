"""WP-05: PRE recon — already in the books (RJ) or the SPV register → no package."""

from __future__ import annotations

import os
from pathlib import Path

import openpyxl
import pytest
from langgraph.checkpoint.memory import MemorySaver

from poarta_contabila.catalog import load_catalog
from poarta_contabila.ingest import IngestDeps, build_ingest_graph, start_payload
from poarta_contabila.jobs import InMemoryJobStore
from poarta_contabila.packages import InMemoryBlobStore, InMemoryPackageStore
from poarta_contabila.recon.numbers import match_level, normalize
from poarta_contabila.recon.pre import (
    InMemoryReconStore,
    PreResult,
    Witnesses,
    check_profile,
    make_pre_check,
)
from poarta_contabila.sinks.exports import ExportError, ExportEye, read_firm_cui, read_saga_rj
from poarta_contabila.sinks.saga_eye import FakeSagaEye
from poarta_contabila.sinks.spv_register import read_spv_register, register_invoices
from poarta_contabila.triage import EmitDecision, Pack
from poarta_contabila.types import (
    CanonicalDocument,
    JobRecord,
    PartnerRef,
    SinkDoc,
    SourceRef,
    TenantRef,
    Totals,
)

SINK = Path(__file__).resolve().parents[1] / "fixtures" / "sink"
TENANT = "1000009"
ALL = ("exact", "alnum", "digits_core")


@pytest.fixture(scope="module")
def cat():
    return load_catalog()


@pytest.fixture(scope="module")
def saga_lines():
    return read_saga_rj(SINK / "saga_rj.xls")


@pytest.fixture(scope="module")
def register():
    return register_invoices(read_spv_register(SINK / "spv_register.xlsx"))


def _eye(lines, **kw):
    kw.setdefault("cui", read_firm_cui(SINK / "saga_rj.xls"))
    return ExportEye(product="saga", lines=lines, **kw)


def _job(period="2026-09") -> JobRecord:
    return JobRecord(
        job_id="job-1",
        tenant=TenantRef(cui=TENANT, saga_firm_folder="0001"),
        period=period,
        status="bound",
    )


def _doc(
    number,
    day,
    gross,
    *,
    doc_class="intrare",
    partner_cui=None,
    partner="FURNIZOR TEST SRL",
    storno=False,
):
    return CanonicalDocument(
        job_id="job-1",
        tenant=TenantRef(cui=TENANT, saga_firm_folder="0001"),
        period=day[:7],
        doc_class=doc_class,
        number=number,
        date=day,
        partner=PartnerRef(
            cui=partner_cui,
            name=partner,
            role="supplier" if "intrare" in doc_class else "customer",
        ),
        totals=Totals(net=gross, vat="0.00", gross=gross),
        lines=[],
        is_storno=storno,
        source=SourceRef(
            kind="ubl_spv", bucket_key="k", content_type="application/xml", source_hash="0" * 64
        ),
    )


def _check(cat, eye, register=(), *, store=None, job=None, doc=None):
    pre = make_pre_check(cat, lambda job: Witnesses(eye=eye, register=list(register)), store=store)
    return pre(job or _job(), doc, fiscal_class="ro_efactura")


# ----- numbers -----


def test_number_levels():
    assert normalize(" ab 0058 ", "exact") == "AB 0058"
    assert normalize("AB-0058", "alnum") == "AB0058"
    assert normalize("AB0058", "digits_core") == "58"
    assert normalize("ABC", "digits_core") is None
    assert match_level("AB 0058", "ab0058", ALL) == "alnum"
    assert match_level("AB0058", "58", ALL) == "digits_core"
    assert match_level("AB0058", "58", ("exact", "alnum")) is None
    assert match_level("FX-101", "FX-101", ALL) == "exact"


# ----- SPV register -----


def test_register_rows_and_invoices(register):
    by = {i.number: i for i in register}
    assert by["1427"].posted and by["1427"].gross == "1210.00"
    assert not by["AB 0058"].posted
    assert by["F 77"].gross == "176.50" and by["F 77"].ref == "register:6,7"  # two VAT rows


def _register_copy(tmp_path, row, col, value):
    wb = openpyxl.load_workbook(SINK / "spv_register.xlsx")
    wb.active.cell(row, col, value)
    path = tmp_path / "reg.xlsx"
    wb.save(path)
    return path


@pytest.mark.parametrize(
    ("row", "col", "value", "match"),
    [
        (4, 10, "Anulat", "unknown status"),
        (4, 9, 1211.00, "gross"),
        (4, 8, 200.00, "gross"),
        (5, 1, "ALTA FIRMA SRL", "more than one company"),
        (7, 3, "2026-09-13", "mixed"),
    ],
)
def test_register_is_checked(tmp_path, row, col, value, match):
    with pytest.raises(ExportError, match=match):
        read_spv_register(_register_copy(tmp_path, row, col, value))


def test_register_vat_must_follow_the_rate(tmp_path):
    wb = openpyxl.load_workbook(SINK / "spv_register.xlsx")
    ws = wb.active
    ws.cell(4, 8, 190.00)
    ws.cell(4, 9, 1190.00)
    path = tmp_path / "reg.xlsx"
    wb.save(path)
    with pytest.raises(ExportError, match="is not 21%"):
        read_spv_register(path)


# ----- profile -----


def test_profiles_carry_the_number_ladder(cat):
    profile = cat.recon_profiles["pre_doc_nr_date"]
    assert profile["number_match"] == list(ALL)
    assert "[de confirmat]" in profile["number_match_note"]
    with pytest.raises(ValueError, match="number_match"):
        check_profile({**profile, "number_match": ["fuzzy"]})
    with pytest.raises(ValueError, match="no number_match"):
        check_profile({**profile, "number_match": None})


# ----- verdicts -----


def test_matching_number_and_date_in_the_journal_is_already_posted(cat, saga_lines):
    eye = _eye(saga_lines, partner_cuis={"401.00010": "40000000"})
    out = _check(cat, eye, doc=_doc("1427", "2026-09-03", "1210.00", partner_cui="40000000"))
    assert (out.verdict, out.witness, out.level) == ("already_posted", "registru_jurnal", "exact")
    assert out.profile_id == "pre_doc_nr_date" and out.hits == ["saga:Intrari:1427:2026-09-03"]


def test_series_typed_differently_matches_at_alnum(cat, saga_lines):
    out = _check(cat, _eye(saga_lines), doc=_doc("AB 0058", "2026-09-10", "167.06"))
    assert (out.verdict, out.level) == ("already_posted", "alnum")


def test_digits_core_needs_the_same_partner_cui(cat, saga_lines):
    doc = _doc("58", "2026-09-10", "167.06", partner_cui="20000005")
    unknown = _check(cat, _eye(saga_lines), doc=doc)
    assert unknown.verdict == "ambiguous" and unknown.near == ["saga:Intrari:AB0058:2026-09-10"]
    known = _check(cat, _eye(saga_lines, partner_cuis={"401.00002": "20000005"}), doc=doc)
    assert (known.verdict, known.level) == ("already_posted", "digits_core")


def test_number_with_another_date_or_gross_is_ambiguous(cat, saga_lines):
    late = _check(cat, _eye(saga_lines), doc=_doc("1427", "2026-09-04", "1210.00"))
    off = _check(cat, _eye(saga_lines), doc=_doc("1427", "2026-09-03", "1250.00"))
    assert late.verdict == off.verdict == "ambiguous"
    assert "close" in late.reason


def test_same_number_from_another_supplier_is_not_a_match(cat, saga_lines):
    eye = _eye(saga_lines, partner_cuis={"401.00010": "40000000"})
    out = _check(cat, eye, doc=_doc("1427", "2026-09-03", "1210.00", partner_cui="30000002"))
    assert out.verdict == "absent"


def test_nothing_close_in_a_covered_month_is_absent(cat, saga_lines):
    out = _check(cat, _eye(saga_lines), doc=_doc("X-9", "2026-09-20", "50.00"))
    assert out.verdict == "absent" and out.witness is None


def test_sales_are_matched_against_sales_only(cat, saga_lines):
    sale = _check(
        cat, _eye(saga_lines), doc=_doc("FX-101", "2026-09-10", "182.11", doc_class="iesire")
    )
    assert sale.verdict == "already_posted"
    other_side = _check(
        cat, _eye(saga_lines), doc=_doc("1427", "2026-09-03", "1210.00", doc_class="iesire")
    )
    assert other_side.verdict == "absent"


def test_no_export_of_the_books_never_says_absent(cat, saga_lines):
    doc = _doc("X-9", "2026-09-20", "50.00")
    assert "need_rj_export" in _check(cat, FakeSagaEye(), doc=doc).reason
    other_firm = _check(cat, _eye(saga_lines, cui="40000000"), doc=doc)
    assert other_firm.verdict == "ambiguous" and "need_rj_export" in other_firm.reason


def test_every_month_from_the_document_to_the_job_must_be_covered(cat, saga_lines):
    doc = _doc("X-9", "2026-08-20", "50.00")
    gap = _check(cat, _eye(saga_lines), doc=doc)
    assert gap.verdict == "ambiguous" and "2026-08" in gap.reason
    full = _check(cat, _eye(saga_lines, periods=["2026-08", "2026-09"]), doc=doc)
    assert full.verdict == "absent"


def test_posted_late_in_a_later_month_is_found(cat, saga_lines):
    # an August invoice the books carry in September: same number/gross, other date → ask
    doc = _doc("1427", "2026-08-28", "1210.00")
    out = _check(cat, _eye(saga_lines, periods=["2026-08", "2026-09"]), doc=doc)
    assert out.verdict == "ambiguous" and out.near == ["saga:Intrari:1427:2026-09-03"]


def test_register_alone_can_say_posted(cat, register):
    empty = ExportEye(product="saga", lines=[], cui=TENANT, periods=["2026-09"])
    doc = _doc("F77", "2026-09-12", "176.50", partner="Al Treilea S.R.L.")
    out = _check(cat, empty, register, doc=doc)
    assert (out.verdict, out.witness) == ("already_posted", "spv_register")
    assert "registru jurnal shows no line" in out.reason


def test_register_still_to_post_is_not_evidence(cat, register):
    empty = ExportEye(product="saga", lines=[], cui=TENANT, periods=["2026-09"])
    out = _check(cat, empty, register, doc=_doc("AB 0058", "2026-09-10", "167.06"))
    assert out.verdict == "absent"


def test_register_is_not_read_for_sales(cat, register):
    empty = ExportEye(product="saga", lines=[], cui=TENANT, periods=["2026-09"])
    doc = _doc("1427", "2026-09-03", "1210.00", doc_class="iesire")
    assert _check(cat, empty, register, doc=doc).verdict == "absent"


class _TwoCopies(FakeSagaEye):
    def covers(self, cui, period):
        return True

    def documents(self, cui, period):
        return [
            SinkDoc(
                saga_key=f"saga:Intrari:1427:{n}",
                doc_class="intrare",
                number="1427",
                date="2026-09-03",
                partner_cui=None,
                gross="1210.00",
                net="1000.00",
                vat="210.00",
                validated=True,
            )
            for n in (1, 2)
        ]


def test_two_hits_are_ambiguous(cat):
    out = _check(cat, _TwoCopies(), doc=_doc("1427", "2026-09-03", "1210.00"))
    assert out.verdict == "ambiguous" and len(out.hits) == 2


def test_storno_is_checked_on_its_own_pre_articol(cat, saga_lines):
    """WP-73 G4 (owner, 2026-10-03): recon_pre_storno; before it a storno always asked."""
    doc = _doc("FX-120", "2026-09-20", "-121.00", doc_class="storn_iesire", storno=True)
    store = InMemoryReconStore()
    out = _check(cat, _eye(saga_lines), store=store, doc=doc)
    assert out.profile_id == "pre_doc_nr_date" and out.verdict == "absent"
    assert len(store.rows) == 1


def test_verdict_is_stored_once_per_sink_snapshot(cat, saga_lines):
    store = InMemoryReconStore()
    doc = _doc("X-9", "2026-09-20", "50.00")
    first = _check(cat, _eye(saga_lines), store=store, doc=doc)
    again = _check(cat, _eye(saga_lines), store=store, doc=doc)
    assert first == again and len(store.rows) == 1
    wider = _check(cat, _eye(saga_lines, periods=["2026-08", "2026-09"]), store=store, doc=doc)
    assert wider.snapshot_id == first.snapshot_id  # August is not read for a September doc
    newer = _check(cat, _eye(saga_lines[:2]), store=store, doc=doc)  # other books, new verdict
    assert newer.snapshot_id != first.snapshot_id and len(store.rows) == 2


# ----- ingest -----


def _ingest(cat, eye, register=()):
    jobs = InMemoryJobStore()
    blobs = InMemoryBlobStore()
    deps = IngestDeps(
        catalog=cat,
        jobs=jobs,
        packages=InMemoryPackageStore(),
        blobs=blobs,
        pre_check=make_pre_check(cat, lambda job: Witnesses(eye=eye, register=list(register))),
        judge=lambda doc, articol: {"accounts_ok": True, "risk": "low", "needs_human": False},
        tenant_name=lambda cui: "FIRMA TEST SRL",
    )
    return deps, build_ingest_graph(deps, checkpointer=MemorySaver())


def _emit(jobs, doc):
    pack = Pack(
        tenant_cui=TENANT,
        saga_firm_folder="0001",
        period="2026-09",
        source_hash=doc.source.source_hash,
        source_doc_id="ro_efactura_ubl",
        kinds=["ubl_spv"],
        our_role="inbound",
        counterparty_cui=doc.partner.cui,
    )
    job = jobs.emit(pack, EmitDecision(emit=True, job_kind="job_ro_efactura", aisle="x")).job
    return job, doc.model_copy(update={"job_id": job.job_id})


def test_ingest_stops_at_already_in_sink_and_never_packages(cat, saga_lines):
    deps, graph = _ingest(cat, _eye(saga_lines))
    job, doc = _emit(deps.jobs, _doc("AB 0058", "2026-09-10", "167.06"))
    cfg = {"configurable": {"thread_id": f"job:{job.job_id}"}}
    out = graph.invoke(start_payload(job, doc, source_doc_id="ro_efactura_ubl"), cfg)
    assert out["status"] == "already_in_sink" and out["pre"]["level"] == "alnum"
    assert deps.jobs.get(job.job_id).status == "already_in_sink"
    assert deps.blobs.puts == 0


def test_ingest_without_books_needs_a_human_with_the_reason(cat):
    deps, graph = _ingest(cat, FakeSagaEye())
    job, doc = _emit(deps.jobs, _doc("X-9", "2026-09-20", "50.00"))
    cfg = {"configurable": {"thread_id": f"job:{job.job_id}"}}
    out = graph.invoke(start_payload(job, doc, source_doc_id="ro_efactura_ubl"), cfg)
    assert out["status"] == "needs_human"
    assert "need_rj_export" in deps.jobs.get(job.job_id).error
    assert deps.blobs.puts == 0


def test_postgres_recon_store_keeps_the_first_verdict(cat):
    dsn = os.environ.get("POARTA_TEST_DSN")
    if not dsn:
        pytest.skip("set POARTA_TEST_DSN to a scratch Postgres")
    from poarta_contabila.jobs import PostgresJobStore
    from poarta_contabila.recon.pre import PostgresReconStore

    jobs = PostgresJobStore(dsn, reset=True)
    job, _ = _emit(jobs, _doc("X-9", "2026-09-20", "50.00"))
    store = PostgresReconStore(dsn)
    first = PreResult(verdict="absent", reason="a", profile_id="pre_doc_nr_date", snapshot_id="s1")
    second = first.model_copy(update={"verdict": "ambiguous", "reason": "b"})
    assert store.put_once(job.job_id, "pre", first) == first
    assert store.put_once(job.job_id, "pre", second) == first


def test_a_credit_note_has_its_own_pre_articol():
    """WP-73 G4: recon_pre_storno takes a storno on the number / date profile."""
    from poarta_contabila.catalog import load_catalog
    from poarta_contabila.flux import MatchContext, match_articole

    cat = load_catalog()
    for storno, want in ((True, ["recon_pre_storno"]), (False, ["recon_pre_standard"])):
        ctx = MatchContext(
            graph_id="reconcile_sink",
            fiscal_class="ro_efactura",
            our_role="outbound",
            is_storno=storno,
            stage="pre",
        )
        assert match_articole(cat, ctx) == want
    assert cat.articole["recon_pre_storno"]["profile_id"] == "pre_doc_nr_date"
