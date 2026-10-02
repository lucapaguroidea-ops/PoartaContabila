"""A2: witness export readers (SAGA C and NextUp) and the export-backed eye."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from poarta_contabila.sinks.exports import (
    ExportError,
    ExportEye,
    read_nextup_balanta,
    read_nextup_rj,
    read_saga_balanta,
    read_saga_rj,
)
from poarta_contabila.sinks.saga_eye import SagaEye

SINK = Path(__file__).resolve().parents[1] / "fixtures" / "sink"


@pytest.fixture(scope="module")
def saga_lines():
    return read_saga_rj(SINK / "saga_rj.xls")


@pytest.fixture(scope="module")
def nextup_lines():
    return read_nextup_rj(SINK / "nextup_rj.xlsx")


# ----- SAGA journal -----


def test_saga_analytic_codes_follow_the_cell_format(saga_lines):
    accounts = {line.credit for line in saga_lines} | {line.debit for line in saga_lines}
    assert "401.00010" in accounts  # stored as the number 401.0001
    assert "401.0001" not in accounts
    assert "5121.090110" in accounts  # 6-decimal analytic
    assert "628" in accounts and "628.0" not in accounts


def test_saga_skips_totals_blank_rows_and_footer(saga_lines):
    assert all(not line.debit.startswith("Total") for line in saga_lines)
    assert [line.seq for line in saga_lines if line.seq][:3] == ["1", "2", "3"]


def test_saga_dates_and_numbers(saga_lines):
    first = saga_lines[0]
    assert first.date == "2026-09-03"
    assert first.doc_number == "1427"  # numeric cell, integer format
    assert first.amount == "1000.00"
    assert {line.doc_number for line in saga_lines if line.journal == "Iesiri"} == {"FX-101"}


def test_saga_compound_entries_expand_into_pairs(saga_lines):
    closing = [line for line in saga_lines if line.journal == "Inchidere"]
    pairs = {(line.debit, line.credit, line.amount) for line in closing}
    assert pairs == {
        ("121", "628", "1000.00"),
        ("121", "605", "150.50"),
        ("704", "121", "150.50"),
    }
    assert not [line for line in closing if "%" in (line.debit, line.credit)]


def test_saga_compound_that_does_not_sum_is_refused(tmp_path):
    import xlrd  # noqa: F401  (format reader present)
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet")
    for c, v in enumerate(
        ["crt.", "Data", "Explicatie", "Nr. doc", "debitor", "creditor", "Debit", "Credit", "Tip"]
    ):
        ws.write(7, c, v)
    ws.write(8, 0, 1)
    ws.write(8, 1, 46295)
    ws.write(8, 4, 121)
    ws.write(8, 5, "%")
    ws.write(8, 6, 100.0)
    ws.write(8, 7, 100.0)
    ws.write(8, 8, "Inchidere")
    ws.write(9, 1, 46295)
    ws.write(9, 5, 628)
    ws.write(9, 6, 90.0)
    ws.write(9, 7, 90.0)
    ws.write(9, 8, "Inchidere")
    path = tmp_path / "bad.xls"
    wb.save(str(path))
    with pytest.raises(ExportError, match="compound"):
        read_saga_rj(path)


# ----- SAGA balance -----


def test_saga_balance_columns_and_codes():
    rows = read_saga_balanta(SINK / "saga_balanta.xlsx")
    by = {}
    for r in rows:
        by.setdefault(r.account, r)
    assert by["401.00010"].closing_credit == "0.00"
    assert by["401.00002"].closing_credit == "166.06"
    assert by["5121.090110"].opening_debit == "5000.00"
    assert by["5121"].closing_debit == "3790.00"
    assert len([r for r in rows if r.account == "1012"]) == 2


# ----- NextUp -----


def test_nextup_journal(nextup_lines):
    assert len(nextup_lines) == 7  # the trailing total row is not a line
    first = nextup_lines[0]
    assert (first.journal, first.doc_number, first.date) == ("JC", "AB0058", "2026-09-10")
    assert first.credit == "401_ALTFURNIZOR" and first.partner == "ALT FURNIZOR SRL"
    assert [line.amount for line in nextup_lines if line.doc_number == "FX-90"] == ["-20.00"]


def test_nextup_auxiliary_accounts_fail_closed(tmp_path):
    wb = openpyxl.load_workbook(SINK / "nextup_rj.xlsx")
    wb.active.cell(2, 8, "AUX1")
    path = tmp_path / "aux.xlsx"
    wb.save(path)
    with pytest.raises(ExportError, match="auxiliar"):
        read_nextup_rj(path)


def test_nextup_balance_gives_partner_cifs():
    rows = read_nextup_balanta(SINK / "nextup_balanta.xlsx")
    by = {r.account: r for r in rows}
    assert by["401_ALTFURNIZOR"].partner_cui == "20000005"  # RO prefix stripped
    assert by["4111_CLIENTTEST"].partner_cui == "30000002"
    assert by["401XVENDOR"].partner_cui is None  # not a valid RO CUI: kept out of keys
    assert by["401"].closing_credit == "167.06"
    assert by["401"].opening_debit is None  # NextUp has no opening columns


def test_wrong_layout_is_refused(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Cont", "Ceva"])
    path = tmp_path / "x.xlsx"
    wb.save(path)
    with pytest.raises(ExportError):
        read_nextup_rj(path)


# ----- eye -----


def test_saga_eye_documents(saga_lines):
    eye = ExportEye(
        product="saga",
        lines=saga_lines,
        balance=read_saga_balanta(SINK / "saga_balanta.xlsx"),
        partner_cuis={"401.00010": "40000000"},
    )
    assert isinstance(eye, SagaEye)
    docs = {d.number: d for d in eye.documents("1000009", "2026-09")}
    assert set(docs) == {"1427", "AB0058", "FX-101", "1"}  # "1": the bank payment (WP-13)
    assert (docs["1"].doc_class, docs["1"].gross, docs["1"].analytic) == (
        "plata",
        "1210.00",
        "5121.090110",
    )
    purchase = docs["1427"]
    assert (purchase.doc_class, purchase.gross, purchase.vat, purchase.net) == (
        "intrare",
        "1210.00",
        "210.00",
        "1000.00",
    )
    assert purchase.analytic == "401.00010" and purchase.partner_cui == "40000000"
    assert docs["AB0058"].gross == "167.06"
    sale = docs["FX-101"]
    assert (sale.doc_class, sale.gross, sale.vat) == ("iesire", "182.11", "31.61")
    assert eye.documents("1000009", "2026-08") == []


def test_nextup_eye_uses_balance_cifs(nextup_lines):
    eye = ExportEye(
        product="nextup",
        lines=nextup_lines,
        balance=read_nextup_balanta(SINK / "nextup_balanta.xlsx"),
    )
    docs = {d.number: d for d in eye.documents("1000009", "2026-09")}
    assert docs["AB0058"].partner_cui == "20000005"
    assert docs["FX-101"].partner_cui == "30000002"
    assert docs["FX-90"].gross == "-20.00"
    assert eye.solduri("1000009", "2026-09")["401"] == {"debit": "0.00", "credit": "167.06"}
    analytic = eye.analytic("1000009", "2026-09", "401")
    assert set(analytic) == {"401_ALTFURNIZOR", "401XVENDOR"}
    assert "4111_CLIENTTEST" not in analytic


def test_vat_on_analytics_counts_4428_tp_ti():
    """R1: SAGA books TVA la încasare on 4428.TP / 4428.TI; the VAT of such an invoice is read."""
    from poarta_contabila.sinks.exports import ExportEye, SinkLine

    def ln(row, journal, number, debit, credit, amount):
        return SinkLine(
            product="saga",
            row=row,
            seq=None,
            date="2026-09-10",
            journal=journal,
            doc_number=number,
            explanation="",
            debit=debit,
            credit=credit,
            amount=amount,
        )

    lines = [
        ln(1, "Intrari", "AB 7", "628", "401.00001", "100.00"),
        ln(2, "Intrari", "AB 7", "4428.TP", "401.00001", "21.00"),
        ln(3, "Iesiri", "FX 8", "4111.00002", "704", "200.00"),
        ln(4, "Iesiri", "FX 8", "4111.00002", "4428.TI", "42.00"),
    ]
    eye = ExportEye(product="saga", lines=lines, cui="1000009", periods=["2026-09"])
    docs = {d.number: d for d in eye.documents("1000009", "2026-09")}
    assert (docs["AB 7"].gross, docs["AB 7"].net, docs["AB 7"].vat) == ("121.00", "100.00", "21.00")
    assert (docs["FX 8"].gross, docs["FX 8"].net, docs["FX 8"].vat) == ("242.00", "200.00", "42.00")
