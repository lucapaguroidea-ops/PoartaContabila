"""The books a firm's documents leave in SAGA (WP-70), and the exports that witness them.

:class:`Book` posts each document the way SAGA posts it — not the way this system expects,
so a gap between the two shows up when the month is run:

- a purchase: the net on the item's expense account against the supplier's analytic; VAT on
  4426 (exigibilitate la livrare), on 4428 (TVA la încasare), or into the cost for a
  neplătitor; reverse charge from abroad for a payer is 4426 = 4427 (autolichidare); for a
  neplătitor it is not generated (WP-D3 open: 4423 vs 446x);
- a sale: the customer's analytic against the income account, VAT on 4427 (4428 la încasare);
- a credit note: the same accounts, negative amounts;
- a bank line: the partner's analytic against the bank analytic, one journal line per invoice
  it settles (TVA la încasare also moves the paid share 4428 → 4426 / 4427); fees on 627,
  salaries on 421;
- a bon: its cost (and VAT on 4426 only for a payer whose CUI is on it) against 542 when it
  sits in an expense report, else 5311;
- a stat de salarii: 641 / 646 against 421, 431x, 444, 436.

Named defects change the books after posting (:meth:`Book.missing`, …). The SAGA journal
types ``Diverse``, ``Casa`` and ``Salarii`` are written for notes, cash and payroll; only
``Intrari`` / ``Iesiri`` / ``Banca`` are read as documents (``[de confirmat]`` on a real book).

Renderers write the shapes of ``fixtures/sink/`` only: SAGA registru jurnal (.xls), balanță
(.xlsx), jurnal de cumpărări / vânzări (.xls); NextUp registru jurnal and balanță (.xlsx);
the SPV register (.xlsx). Workbook timestamps are fixed: the same book gives the same bytes.
SAGA's .xls files are written with ``xlwt`` (a dev dependency, imported only when rendered).
"""

from __future__ import annotations

import io
import zipfile
from collections import defaultdict
from dataclasses import dataclass, replace
from datetime import date, datetime
from decimal import Decimal
from typing import Literal

from poarta_contabila.synthetic.docs import (
    ZIP_TIME,
    BankLine,
    Bon,
    Invoice,
    Payroll,
    vat_of,
)
from poarta_contabila.synthetic.firms import Firm, Party

FX_EUR = Decimal("4.9750")  # an invented fixed rate for documents in EUR
FIXED = datetime(2026, 10, 1)
_TIP_ORDER = {"Intrari": 0, "Iesiri": 1, "Banca": 2, "Casa": 3, "Diverse": 4, "Salarii": 5}


@dataclass(frozen=True)
class Entry:
    """One journal line: debit account, credit account, one amount (cents; negative = storno)."""

    ref: str  # the document it comes from
    day: str
    tip: str  # SAGA journal type
    number: str | None
    expl: str
    debit: str
    credit: str
    amount: int
    kind: Literal["net", "vat", "bank", "other"] = "other"
    rate: int | None = None
    partner: Party | None = None


Doc = Invoice | BankLine | Bon | Payroll


def _ron(inv: Invoice, cents: int) -> int:
    if inv.currency == "RON":
        return cents
    return int((Decimal(cents) * FX_EUR).quantize(Decimal(1)))


def _nb(number: str | None) -> str | None:
    """SAGA keeps a document number without spaces (``AB 0058`` → ``AB0058``)."""
    return number.replace(" ", "") if number else number


class Book:
    """A firm's journal in SAGA, built document by document, months in any order."""

    def __init__(self, firm: Firm) -> None:
        self.firm = firm
        self.entries: list[Entry] = []
        self.invoices: dict[str, Invoice] = {}
        self.skew: dict[str, int] = {}  # balanță: synthetic row off its analytics (cents)

    # ----- posting -----

    def post(self, *docs: Doc, via: str | None = None) -> Book:
        """Post *docs* as SAGA would; ``via`` = the account a bon is paid from (542 / 5311)."""
        for doc in docs:
            if isinstance(doc, Invoice):
                self.invoices[doc.ref] = doc
                self.entries += self._invoice(doc)
            elif isinstance(doc, BankLine):
                self.entries += self._bank(doc)
            elif isinstance(doc, Bon):
                self.entries += self._bon(doc, via or self.firm.cash_account)
            elif isinstance(doc, Payroll):
                self.entries += self._payroll(doc)
            else:
                raise TypeError(f"cannot post {type(doc).__name__}")
        return self

    def _invoice(self, inv: Invoice) -> list[Entry]:
        f, p, s = self.firm, inv.partner, (-1 if inv.is_credit else 1)
        num, out = _nb(inv.number), []
        if inv.side == "in":
            tip, expl = "Intrari", f"Intrare {p.name}"
            for it in inv.items:
                out.append(
                    Entry(
                        inv.ref,
                        inv.issued,
                        tip,
                        num,
                        expl,
                        it.account,
                        p.analytic,
                        s * _ron(inv, it.net),
                        "net",
                        it.rate,
                        p,
                    )
                )
            for (cat, rate), (taxable, vat) in inv.groups().items():
                if cat == "AE":
                    if f.vat_payer:  # autolichidare: 4426 = 4427 on the reverse-charged base
                        rc = vat_of(_ron(inv, taxable), 21)
                        out.append(
                            Entry(
                                inv.ref,
                                inv.issued,
                                tip,
                                num,
                                f"TVA taxare inversa {p.name}",
                                "4426",
                                "4427",
                                s * rc,
                                "vat",
                                21,
                                p,
                            )
                        )
                    continue  # neplătitor: WP-D3 (4423 vs 446x) is open, nothing is invented
                if not vat:
                    continue
                if not f.vat_payer:
                    cost = next(it.account for it in inv.items if it.rate == rate)
                    debit = cost  # a neplătitor carries the VAT in the cost
                else:
                    debit = "4428" if f.la_incasare else "4426"
                out.append(
                    Entry(
                        inv.ref,
                        inv.issued,
                        tip,
                        num,
                        f"TVA {rate} {p.name}",
                        debit,
                        p.analytic,
                        s * _ron(inv, vat),
                        "vat",
                        rate,
                        p,
                    )
                )
        else:
            tip, expl = "Iesiri", f"Iesire {p.name}"
            for it in inv.items:
                out.append(
                    Entry(
                        inv.ref,
                        inv.issued,
                        tip,
                        num,
                        expl,
                        p.analytic,
                        it.account,
                        s * _ron(inv, it.net),
                        "net",
                        it.rate,
                        p,
                    )
                )
            for (_, rate), (_, vat) in inv.groups().items():
                if vat:
                    credit = "4428" if f.la_incasare else "4427"
                    out.append(
                        Entry(
                            inv.ref,
                            inv.issued,
                            tip,
                            num,
                            f"TVA {rate} {p.name}",
                            p.analytic,
                            credit,
                            s * _ron(inv, vat),
                            "vat",
                            rate,
                            p,
                        )
                    )
        return out

    def _bank(self, ln: BankLine) -> list[Entry]:
        f, bank, num, out = self.firm, self.firm.bank_account, ln.reference, []
        ref = ln.reference or ln.description
        if ln.kind == "fee":
            return [
                Entry(ref, ln.day, "Banca", num, ln.description, "627", bank, ln.amount, "bank")
            ]
        if ln.kind == "salary":
            return [
                Entry(ref, ln.day, "Banca", num, ln.description, "421", bank, ln.amount, "bank")
            ]
        if ln.kind == "other" or not ln.settles:
            debit, credit = ("461", bank) if ln.side == "debit" else (bank, "461")
            return [
                Entry(ref, ln.day, "Banca", num, ln.description, debit, credit, ln.amount, "bank")
            ]
        for inv_ref, amount in ln.settles:
            inv = self.invoices.get(inv_ref)
            p = inv.partner if inv is not None else ln.partner
            if p is None:
                raise ValueError(f"bank line {num}: {inv_ref!r} is not posted in this book")
            if ln.side == "debit":
                out.append(
                    Entry(
                        ref,
                        ln.day,
                        "Banca",
                        num,
                        f"Achit. {p.name} {inv.number if inv else ''}".strip(),
                        p.analytic,
                        bank,
                        amount,
                        "bank",
                        partner=p,
                    )
                )
            else:
                out.append(
                    Entry(
                        ref,
                        ln.day,
                        "Banca",
                        num,
                        f"Incasare {p.name} {inv.number if inv else ''}".strip(),
                        bank,
                        p.analytic,
                        amount,
                        "bank",
                        partner=p,
                    )
                )
            if f.la_incasare and inv is not None and inv.vat:
                share = int((Decimal(amount) * inv.vat / inv.gross).quantize(Decimal(1)))
                debit, credit = ("4426", "4428") if ln.side == "debit" else ("4428", "4427")
                out.append(
                    Entry(
                        ref,
                        ln.day,
                        "Banca",
                        num,
                        f"TVA exigibila {p.name}",
                        debit,
                        credit,
                        share,
                        "vat",
                        partner=p,
                    )
                )
        return out

    def _bon(self, bon: Bon, via: str) -> list[Entry]:
        tip = "Diverse" if via.startswith("542") else "Casa"
        num, out = f"BF{bon.number}", []
        deduct = self.firm.vat_payer and bon.our_cui is not None
        for it in bon.items:
            vat = vat_of(it.net, it.rate)
            out.append(
                Entry(
                    bon.ref,
                    bon.day,
                    tip,
                    num,
                    f"Bon {bon.merchant.name}",
                    it.account,
                    via,
                    it.net if deduct else it.net + vat,
                    "net",
                    it.rate,
                )
            )
            if deduct and vat:
                out.append(
                    Entry(
                        bon.ref,
                        bon.day,
                        tip,
                        num,
                        f"TVA bon {bon.merchant.name}",
                        "4426",
                        via,
                        vat,
                        "vat",
                        it.rate,
                    )
                )
        return out

    def _payroll(self, pay: Payroll) -> list[Entry]:
        day = f"{pay.period}-{date.fromisoformat(pay.day).day:02d}"
        rows = (
            ("641", "421", pay.gross, "Salarii brute"),
            ("421", "4315", pay.cas, "CAS"),
            ("421", "4316", pay.cass, "CASS"),
            ("421", "444", pay.tax, "Impozit salarii"),
            ("646", "436", pay.cam, "CAM"),
        )
        return [
            Entry(pay.ref, day, "Salarii", f"SAL{pay.period[2:4]}{pay.period[5:]}", expl, d, c, a)
            for d, c, a, expl in rows
        ]

    # ----- named defects -----

    def _of(self, ref: str) -> list[int]:
        idx = [i for i, e in enumerate(self.entries) if e.ref == ref]
        if not idx:
            raise KeyError(f"{ref!r} is not in the book")
        return idx

    def missing(self, ref: str) -> Book:
        """missing from the books: the document was never posted."""
        self._of(ref)
        self.entries = [e for e in self.entries if e.ref != ref]
        return self

    def amount_differs(self, ref: str, delta: int) -> Book:
        """the amount in the books differs: the first net line is off by *delta* cents."""
        i = next(i for i in self._of(ref) if self.entries[i].kind in ("net", "bank"))
        e = self.entries[i]
        self.entries[i] = replace(e, amount=e.amount + delta)
        return self

    def vat_differs(self, ref: str, delta: int) -> Book:
        """the VAT in the books differs by *delta* cents (and so the gross)."""
        i = next(i for i in self._of(ref) if self.entries[i].kind == "vat")
        e = self.entries[i]
        self.entries[i] = replace(e, amount=e.amount + delta)
        return self

    def posted_another_way(self, ref: str) -> Book:
        """posted another way: same gross, the VAT not on its account (a payer's purchase:
        into the cost; a neplătitor's purchase: deducted on 4426; a sale: on 4428, not 4427)."""
        idx = self._of(ref)
        changed = False
        for i in idx:
            e = self.entries[i]
            if (
                e.kind != "vat"
                or e.partner is None
                or e.partner.analytic
                not in (
                    e.debit,
                    e.credit,
                )
            ):
                continue
            if e.debit in ("4426", "4428"):
                cost = next(self.entries[j].debit for j in idx if self.entries[j].kind == "net")
                self.entries[i] = replace(e, debit=cost)
            elif e.credit == "4427":
                self.entries[i] = replace(e, credit="4428")
            elif e.credit == e.partner.analytic:  # a neplătitor carried it in the cost
                self.entries[i] = replace(e, debit="4426")
            else:
                continue
            changed = True
        if not changed:
            raise ValueError(f"{ref!r} has no VAT line to post another way")
        return self

    def skew_analytic(self, root: str, delta: int) -> Book:
        """balanță: the synthetic *root* row closes *delta* cents off the sum of its analytics."""
        self.skew[root] = self.skew.get(root, 0) + delta
        return self

    def extra(self, entry: Entry) -> Book:
        """A line in the books with no document here (bank fee, payroll, adjustments …)."""
        self.entries.append(entry)
        return self

    # ----- reading -----

    def periods(self) -> list[str]:
        return sorted({e.day[:7] for e in self.entries})

    def lines(self, periods: list[str] | None = None) -> list[Entry]:
        want = set(periods or self.periods())
        return sorted(
            (e for e in self.entries if e.day[:7] in want),
            key=lambda e: (e.day, _TIP_ORDER.get(e.tip, 9), e.number or "", e.ref),
        )

    def balances(self, period: str) -> dict[str, dict[str, int]]:
        """Per booked account: opening (year), previous months, the period's turnover."""
        year = period[:4]
        out: dict[str, dict[str, int]] = defaultdict(
            lambda: {"open_d": 0, "open_c": 0, "prev_d": 0, "prev_c": 0, "run_d": 0, "run_c": 0}
        )
        for account, cents in self.firm.opening.items():
            out[account]["open_d" if cents > 0 else "open_c"] += abs(cents)
        for e in self.entries:
            if e.day[:4] != year or e.day[:7] > period:
                continue
            col = "run" if e.day[:7] == period else "prev"
            out[e.debit][f"{col}_d"] += e.amount
            out[e.credit][f"{col}_c"] += e.amount
        return dict(out)

    # ----- SAGA renderers -----

    def saga_rj(self, periods: list[str] | None = None) -> bytes:
        """SAGA C *Registru jurnal* (.xls), the layout of ``fixtures/sink/saga_rj.xls``."""
        import xlwt

        wb = xlwt.Workbook()
        ws = wb.add_sheet("Sheet")
        styles: dict[str, xlwt.XFStyle] = {}

        def style(fmt: str) -> xlwt.XFStyle:
            if fmt not in styles:
                styles[fmt] = xlwt.easyxf(num_format_str=fmt)
            return styles[fmt]

        ws.write(0, 0, self.firm.header)
        ws.write(3, 0, "REGISTRU JURNAL")
        for c, v in enumerate(["Nr.", "", "", "", "Cont", "Cont", "", "", ""]):
            ws.write(6, c, v)
        heads = [
            "crt.",
            "Data",
            "Explicatie",
            "Nr. doc",
            "debitor",
            "creditor",
            "Debit",
            "Credit",
            "Tip",
        ]
        for c, v in enumerate(heads):
            ws.write(7, c, v)
        row, seq = 8, 0
        day_total = month_total = 0
        lines = self.lines(periods)
        for i, e in enumerate(lines):
            seq += 1
            ws.write(row, 0, seq, style("###,###,##0"))
            ws.write(row, 1, _serial(e.day), style("m/d/yy"))
            ws.write(row, 2, e.expl)
            if e.number and e.number.isdigit():
                ws.write(row, 3, int(e.number), style("########0"))
            elif e.number:
                ws.write(row, 3, e.number)
            for c, account in ((4, e.debit), (5, e.credit)):
                if "." in account:
                    frac = account.split(".")[1]
                    ws.write(row, c, float(account), style("########0." + "0" * len(frac)))
                else:
                    ws.write(row, c, int(account), style("########0"))
            ws.write(row, 6, e.amount / 100, style("###,###,##0.00"))
            ws.write(row, 7, e.amount / 100, style("###,###,##0.00"))
            ws.write(row, 8, e.tip)
            row += 1
            day_total += e.amount
            month_total += e.amount
            nxt = lines[i + 1] if i + 1 < len(lines) else None
            if nxt is None or nxt.day != e.day:
                ws.write(row, 4, "Total pe")
                ws.write(row, 5, _serial(e.day), style("m/d/yy"))
                ws.write(row, 6, day_total / 100, style("###,###,##0.00"))
                ws.write(row, 7, day_total / 100, style("###,###,##0.00"))
                row, day_total = row + 1, 0
            if nxt is None or nxt.day[:7] != e.day[:7]:
                ws.write(row, 4, "Total luna")
                ws.write(row, 5, int(e.day[5:7]), style("########0"))
                ws.write(row, 6, month_total / 100, style("###,###,##0.00"))
                ws.write(row, 7, month_total / 100, style("###,###,##0.00"))
                row, month_total = row + 2, 0
        ws.write(row + 2, 0, "Pagina 1/1  SAGA C")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def saga_balanta(self, period: str) -> bytes:
        """SAGA C *Balanță de verificare* (.xlsx), the layout of ``saga_balanta.xlsx``."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet"
        ws.cell(1, 1, self.firm.header)
        ws.cell(3, 3, "Balanta de verificare")
        ws.cell(4, 7, datetime(int(period[:4]), 1, 1))
        ws.cell(4, 8, "--")
        ws.cell(4, 9, datetime.fromisoformat(_month_end(period)))
        groups = [
            "Solduri initiale an",
            "Sume precedente",
            "Rulaje perioada",
            "Sume totale",
            "Solduri finale",
        ]
        for i, g in enumerate(groups):
            ws.cell(6, 5 + 2 * i, g)
        ws.cell(7, 1, "Cont")
        ws.cell(7, 2, "Denumirea contului")
        for i in range(5):
            ws.cell(8, 5 + 2 * i, "Debitoare")
            ws.cell(8, 6 + 2 * i, "Creditoare")
        r = 9
        for account, name, vals in self._balance_rows(period):
            c = ws.cell(r, 1)
            if "." in account:
                c.value = float(account)
                c.number_format = "########0." + "0" * len(account.split(".")[1])
            else:
                c.value = int(account)
                c.number_format = "########0"
            ws.cell(r, 2, name)
            for i, v in enumerate(vals):
                ws.cell(r, 5 + i, v / 100).number_format = "###,###,##0.00"
            r += 1
        ws.cell(r + 3, 1, "Pagina 1/1")
        return _stable_xlsx(wb)

    def _balance_rows(self, period: str) -> list[tuple[str, str, list[int]]]:
        bal = self.balances(period)
        names = self._names()
        accounts = set(bal)
        for a in list(accounts):
            if "." in a:
                accounts.add(a.split(".")[0])
        rows = []
        for account in sorted(accounts, key=lambda a: (a.split(".")[0], a)):
            if "." in account or not any(k.startswith(account + ".") for k in bal):
                b = bal.get(account) or {}
            else:  # a synthetic row over its analytics
                b = defaultdict(int)
                for k, v in bal.items():
                    if k == account or k.startswith(account + "."):
                        for col, cents in v.items():
                            b[col] += cents
            rows.append(
                (
                    account,
                    names.get(account, f"CONT {account}"),
                    _columns(b, self.skew.get(account, 0) if "." not in account else 0),
                )
            )
        return rows

    def _names(self) -> dict[str, str]:
        names = {
            "401": "FURNIZORI",
            "4111": "CLIENTI",
            "5121": "CONTURI LA BANCI IN LEI",
            "5311": "CASA IN LEI",
            "542": "AVANSURI DE TREZORERIE",
            "4426": "TVA DEDUCTIBILA",
            "4427": "TVA COLECTATA",
            "4428": "TVA NEEXIGIBILA",
            "1012": "CAPITAL SUBSCRIS VARSAT",
            "117": "REZULTATUL REPORTAT",
            self.firm.bank_account: "BANCA SINTETICA LEI",
        }
        for p in (*self.firm.suppliers, *self.firm.customers, *self.firm.foreign):
            names[p.analytic] = p.name
        return names

    def tva_journal(self, period: str, side: Literal["cumparari", "vanzari"]) -> bytes:
        """SAGA's *Jurnal de cumpărări* / *vânzări* (.xls), the layout of the fixtures; a
        ``Baza 0%`` column is added only when a document has a 0 % base."""
        import xlwt

        tip = "Intrari" if side == "cumparari" else "Iesiri"
        docs: dict[str, dict] = {}
        for e in self.lines([period]):
            if e.tip != tip:
                continue
            d = docs.setdefault(
                e.ref,
                {
                    "day": e.day,
                    "number": e.number,
                    "partner": e.partner,
                    "base": defaultdict(int),
                    "vat": defaultdict(int),
                },
            )
            if e.kind == "net":
                d["base"][e.rate or 0] += e.amount
            elif e.kind == "vat" and e.partner and e.partner.analytic in (e.debit, e.credit):
                d["vat"][e.rate or 0] += e.amount  # the invoice's VAT, not autolichidare
        rates = [21, 11] + ([0] if any(d["base"].get(0) for d in docs.values()) else [])
        wb = xlwt.Workbook()
        ws = wb.add_sheet("Sheet")
        money = xlwt.easyxf(num_format_str="###,###,##0.00")
        dfmt = xlwt.easyxf(num_format_str="dd.mm.yyyy")
        ws.write(0, 0, self.firm.header)
        ws.write(3, 0, "JURNAL DE CUMPARARI" if side == "cumparari" else "JURNAL DE VANZARI")
        ws.write(4, 0, f"Luna {period}")
        head = ["Nr. crt", "Data", "Nr. doc", "Denumire partener", "Cod fiscal", "Total document"]
        for rate in rates:
            head += [f"Baza {rate}%", f"TVA {rate}%"]
        for c, v in enumerate(head):
            ws.write(6, c, v)
        r, total = 7, 0
        for i, d in enumerate(
            sorted(docs.values(), key=lambda d: (d["day"], d["number"] or "")), start=1
        ):
            p: Party | None = d["partner"]
            doc_total = sum(d["base"].values()) + sum(d["vat"].values())
            total += doc_total
            ws.write(r, 0, i)
            ws.write(r, 1, _serial(d["day"]), dfmt)
            ws.write(r, 2, d["number"] or "")
            ws.write(r, 3, p.name if p else "")
            ws.write(r, 4, p.tax_id if p else "")
            ws.write(r, 5, doc_total / 100, money)
            for k, rate in enumerate(rates):
                ws.write(r, 6 + 2 * k, d["base"].get(rate, 0) / 100, money)
                ws.write(r, 7 + 2 * k, d["vat"].get(rate, 0) / 100, money)
            r += 1
        ws.write(r, 0, "Total")
        ws.write(r, 5, total / 100, money)
        ws.write(r + 3, 0, "Pagina 1/1  SAGA C")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ----- NextUp renderers (A2: a read-only eye) -----

    def _nextup_account(self, account: str, partner: Party | None) -> str:
        root = account.split(".")[0]
        if "." in account and partner is not None:
            return f"{root}_{partner.short}"
        return root

    def nextup_rj(self, periods: list[str] | None = None) -> bytes:
        """NextUp *Registru jurnal* (.xlsx), the layout of ``fixtures/sink/nextup_rj.xlsx``."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet"
        ws.append(
            [
                "Nr. crt",
                "Jurnal",
                "Nr. document",
                "Data",
                "Partener",
                "Explicatii",
                "Cont debitor",
                "Cont auxiliar debitor",
                "Cont creditor",
                "Cont auxiliar creditor",
                "Suma",
                "Nr. inregistrare",
            ]
        )
        journal = {"Intrari": "JC", "Iesiri": "JV", "Banca": "JB"}
        notes: dict[tuple[str, str | None, str], int] = {}
        total = 0
        for i, e in enumerate(self.lines(periods), start=1):
            note = notes.setdefault((e.tip, e.number, e.day), len(notes) + 1)
            ws.append(
                [
                    i,
                    journal.get(e.tip, "OD"),
                    e.number,
                    datetime.fromisoformat(e.day),
                    e.partner.name if e.partner else None,
                    e.expl,
                    self._nextup_account(e.debit, e.partner),
                    None,
                    self._nextup_account(e.credit, e.partner),
                    None,
                    e.amount / 100,
                    note,
                ]
            )
            total += e.amount
        ws.append([None] * 10 + [total / 100, None])
        return _stable_xlsx(wb)

    def nextup_balanta(self, period: str) -> bytes:
        """NextUp *Balanță* (.xlsx), the layout of ``fixtures/sink/nextup_balanta.xlsx``."""
        import openpyxl

        partners = {
            p.analytic: p for p in (*self.firm.suppliers, *self.firm.customers, *self.firm.foreign)
        }
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet"
        ws.append(
            [
                "Cont",
                "Titlu cont",
                "Cod partener",
                "Denumire partener",
                "CIF partener",
                "TVA la incasare",
                "Rulaj precedent",
                None,
                "Rulaj curent",
                None,
                "Total",
                None,
                "Sold final",
                None,
            ]
        )
        ws.append([None] * 6 + ["Debit", "Credit"] * 4)
        sums = [0] * 8
        for account, name, vals in self._balance_rows(period):
            # SAGA's columns: SI d/c, prec d/c, run d/c, total d/c, final d/c → NextUp's four
            prev_d, prev_c, run_d, run_c, tot_d, tot_c, fin_d, fin_c = vals[2:]
            p = partners.get(account)
            code = self._nextup_account(account, p)
            row = [
                code,
                name,
                f"{10000 + len(sums)}" if p else None,
                p.name if p else None,
                p.tax_id if p else None,
                ("Da" if self.firm.la_incasare else "Nu") if p else None,
                prev_d / 100,
                prev_c / 100,
                run_d / 100,
                run_c / 100,
                tot_d / 100,
                tot_c / 100,
                fin_d / 100,
                fin_c / 100,
            ]
            ws.append(row)
            if "." not in account:
                sums = [a + b for a, b in zip(sums, vals[2:], strict=True)]
        ws.append([None] * 6 + [v / 100 for v in sums])
        return _stable_xlsx(wb)

    # ----- SPV register -----

    def spv_register(self, invoices: list[Invoice], period: str) -> bytes:
        """The SPV invoice register: one row per purchase and VAT rate; posted when the book
        holds it ("Înregistrat în SAGA"), else "De înregistrat"."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Facturi SPV"
        ws.append(["Registru facturi SPV"])
        ws.append([])
        ws.append(
            [
                "Companie",
                "Trimestru",
                "Data facturii",
                "Numar factura",
                "Numele furnizorului",
                "Suma net per % de TVA",
                "Cota TVA",
                "TVA",
                "Gross",
                "Status",
                "Obs re status",
                "Ordine",
            ]
        )
        posted = {e.ref for e in self.entries}
        quarter = f"T{(int(period[5:]) - 1) // 3 + 1} {period[:4]}"
        order = 0
        for inv in sorted(invoices, key=lambda i: (i.issued, i.number)):
            if inv.side != "in" or inv.foreign or inv.issued[:7] != period:
                continue
            order += 1
            status = "Înregistrat în SAGA" if inv.ref in posted else "De înregistrat"
            s = -1 if inv.is_credit else 1
            for (_, rate), (taxable, vat) in inv.groups().items():
                ws.append(
                    [
                        self.firm.name,
                        quarter,
                        datetime.fromisoformat(inv.issued),
                        inv.number,
                        inv.partner.name,
                        s * taxable / 100,
                        rate / 100,
                        s * vat / 100,
                        s * (taxable + vat) / 100,
                        status,
                        None,
                        order,
                    ]
                )
                r = ws.max_row
                ws.cell(r, 7).number_format = "0%"
                for c in (6, 8, 9):
                    ws.cell(r, c).number_format = "#,##0.00"
        return _stable_xlsx(wb)


def _columns(b, skew: int) -> list[int]:
    od, oc = b.get("open_d", 0), b.get("open_c", 0)
    pd_, pc = od + b.get("prev_d", 0), oc + b.get("prev_c", 0)
    rd, rc = b.get("run_d", 0), b.get("run_c", 0)
    td, tc = pd_ + rd, pc + rc
    net = td - tc + skew
    open_net = od - oc
    return [max(open_net, 0), max(-open_net, 0), pd_, pc, rd, rc, td, tc, max(net, 0), max(-net, 0)]


def _serial(day: str) -> int:
    return (date.fromisoformat(day) - date(1899, 12, 30)).days


def _month_end(period: str) -> str:
    import calendar

    y, m = int(period[:4]), int(period[5:])
    return f"{period}-{calendar.monthrange(y, m)[1]:02d}"


def _stable_xlsx(wb) -> bytes:
    """Save *wb* with fixed document properties and zip timestamps: the same bytes each run."""
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.properties.creator = "synthetic"
    wb.properties.lastModifiedBy = "synthetic"
    raw = io.BytesIO()
    wb.save(raw)
    src = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = _fixed_core(data)
            zf.writestr(zipfile.ZipInfo(info.filename, date_time=ZIP_TIME), data)
    return out.getvalue()


def _fixed_core(data: bytes) -> bytes:
    import re

    stamp = FIXED.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    return re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>" + stamp, data)
